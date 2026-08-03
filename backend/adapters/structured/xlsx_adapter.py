from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.domain.evidence import QualityFlag, RawUnitDraft
from backend.domain.locators import XlsxRowLocator
from backend.domain.sources import Source

from .common import (
    ADAPTER_VERSION,
    ParsedRecord,
    StructuredInspection,
    ambiguous_mapping_exceptions,
    build_structure_profile,
    normalize_header,
    quality_flags,
    raw_hash,
    stable_raw_id,
)


def _first_data_row(sheet) -> int | None:
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if any(value not in (None, "") for value in row):
            return index
    return None


def inspect_xlsx(path: Path, source: Source) -> StructuredInspection:
    payload = path.read_bytes()
    formulas = load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
    cached = load_workbook(io.BytesIO(payload), read_only=False, data_only=True)
    structures = []
    records: list[ParsedRecord] = []
    exceptions: list[dict[str, Any]] = []
    for formula_sheet in formulas.worksheets:
        cached_sheet = cached[formula_sheet.title]
        structure_id = f"xlsx:{formula_sheet.title}"
        hidden = formula_sheet.sheet_state in {"hidden", "veryHidden"}
        header_index = _first_data_row(formula_sheet)
        if header_index is None:
            structures.append(
                build_structure_profile(
                    structure_id=structure_id,
                    name=formula_sheet.title,
                    kind="sheet",
                    rows=[],
                    included=not hidden,
                    hidden=hidden,
                )
            )
            continue
        header_cells = list(formula_sheet.iter_rows(min_row=header_index, max_row=header_index))[0]
        seen: dict[str, int] = {}
        headers = [normalize_header(cell.value, index + 1, seen) for index, cell in enumerate(header_cells)]
        rows: list[dict[str, Any]] = []
        for row_index in range(header_index + 1, formula_sheet.max_row + 1):
            formula_cells = list(formula_sheet.iter_rows(min_row=row_index, max_row=row_index))[0]
            cached_cells = list(cached_sheet.iter_rows(min_row=row_index, max_row=row_index))[0]
            if not any(cell.value not in (None, "") for cell in formula_cells):
                continue
            values: list[Any] = []
            formula_missing_cache = False
            for formula_cell, cached_cell in zip(formula_cells, cached_cells, strict=True):
                if formula_cell.data_type == "f":
                    value = cached_cell.value
                    if value is None:
                        formula_missing_cache = True
                else:
                    value = formula_cell.value
                values.append(value)
            row = dict(zip(headers, values, strict=True))
            rows.append(row)
            cells = [f"{get_column_letter(index + 1)}{row_index}" for index in range(len(headers))]
            locator = XlsxRowLocator(sheet=formula_sheet.title, row=row_index, cells=cells)
            locator_payload = locator.model_dump(mode="json")
            disposition = "excluded" if hidden else "quarantined" if formula_missing_cache else "processed"
            reason = (
                "HIDDEN_SHEET_NOT_INCLUDED"
                if hidden
                else "FORMULA_WITHOUT_CACHED_VALUE"
                if formula_missing_cache
                else "STRUCTURED_RECORD_PREPARED"
            )
            records.append(
                ParsedRecord(
                    raw_unit=RawUnitDraft(
                        raw_unit_id=stable_raw_id(source.source_id, structure_id, locator_payload),
                        parent_raw_unit_id=None,
                        unit_kind="xlsx_row",
                        source_id=source.source_id,
                        structure_id=structure_id,
                        locator=locator,
                        raw_hash=raw_hash(row),
                        adapter_version=ADAPTER_VERSION,
                        quality_flags=(
                            quality_flags(QualityFlag.MISSING_REQUIRED_SOURCE_FIELD)
                            if formula_missing_cache
                            else []
                        ),
                    ),
                    values=row,
                    included=not hidden and not formula_missing_cache,
                    disposition=disposition,
                    reason_code=reason,
                )
            )
            if formula_missing_cache:
                exceptions.append(
                    {
                        "exception_kind": "formula_without_cached_value",
                        "severity": "warning",
                        "title": f"Una formula non ha un valore disponibile in “{formula_sheet.title}”",
                        "explanation": "La riga è stata isolata; le altre righe e gli altri file continuano normalmente.",
                        "payload": {"structure_id": structure_id, "row": row_index, "cells": cells},
                    }
                )
        structure = build_structure_profile(
            structure_id=structure_id,
            name=formula_sheet.title,
            kind="sheet",
            rows=rows,
            included=not hidden,
            hidden=hidden,
        )
        structures.append(structure)
        exceptions.extend(ambiguous_mapping_exceptions(structure))
        if hidden:
            exceptions.append(
                {
                    "exception_kind": "hidden_sheet",
                    "severity": "warning",
                    "title": f"Il foglio “{formula_sheet.title}” è nascosto",
                    "explanation": "È stato inventariato ma non incluso automaticamente. Puoi mantenerlo escluso o includerlo consapevolmente.",
                    "payload": {"structure_id": structure_id, "sheet": formula_sheet.title, "suggested_included": False},
                }
            )
    formulas.close()
    cached.close()
    return StructuredInspection(structures=structures, records=records, exceptions=exceptions)
